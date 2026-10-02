"""Structured, local document extraction without local OCR.

Native formats are parsed locally. Images/scanned PDFs intentionally remain OCR-free
until the user explicitly configures and selects Baidu Cloud "文档解析（PaddleOCR-VL）".
Every extractor emits reading-order blocks so tables never drift to the end of a document.
"""
from __future__ import annotations
import hashlib, re, shutil, subprocess, tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

@dataclass(slots=True)
class PageContent:
    page_number: int
    text: str
    headings: list[str] = field(default_factory=list)

@dataclass(slots=True)
class LayoutBlock:
    order: int
    block_type: str
    text: str
    page_number: int | None = None
    section: str = ""
    bbox: list[float] | None = None
    table: list[list[str]] | None = None
    sheet_name: str | None = None

@dataclass(slots=True)
class ExtractedText:
    text: str
    pages: list[PageContent] = field(default_factory=list)
    blocks: list[LayoutBlock] = field(default_factory=list)
    used_ocr: bool = False
    warning: str = ""
    source_method: str = "native_unknown"
    parser_name: str = ""
    requires_cloud_ocr: bool = False


def file_checksum(path: str | Path) -> str:
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for block in iter(lambda:f.read(1024*1024),b""):
            h.update(block)
    return h.hexdigest()

_HEADING_RE=re.compile(r"^(第[一二三四五六七八九十百千万0-9]+[章节条款编]|[一二三四五六七八九十]+[、．.]|\d+(?:\.\d+)*[、.\s]|[（(][一二三四五六七八九十0-9]+[）)])")
def _find_headings(text:str)->list[str]:
    return [x.strip() for x in text.splitlines() if x.strip() and len(x.strip())<140 and _HEADING_RE.match(x.strip())]

def _cell(value: Any) -> str:
    return str(value if value is not None else "").replace("\n","<br>").strip()

def table_markdown(rows: list[list[Any]]) -> str:
    rows=[[ _cell(c) for c in row] for row in rows]
    rows=[r for r in rows if any(c for c in r)]
    if not rows:return ""
    width=max(len(r) for r in rows)
    rows=[r+[""]*(width-len(r)) for r in rows]
    header=rows[0]
    out=["| "+" | ".join(x.replace("|",r"\|") for x in header)+" |",
         "| "+" | ".join("---" for _ in header)+" |"]
    out += ["| "+" | ".join(x.replace("|",r"\|") for x in row)+" |" for row in rows[1:]]
    return "\n".join(out)

def _plain(path:Path)->ExtractedText:
    raw=path.read_bytes()
    text=""
    for enc in ("utf-8","utf-8-sig","gb18030","latin-1"):
        try: text=raw.decode(enc);break
        except UnicodeDecodeError: continue
    block=LayoutBlock(1,"paragraph",text,page_number=1)
    return ExtractedText(text=text,pages=[PageContent(1,text,_find_headings(text))],blocks=[block],source_method="native_txt",parser_name="plain-text")

def _iter_docx_blocks(document) -> Iterator[tuple[str, Any]]:
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in document.element.body.iterchildren():
        if child.tag==qn("w:p"): yield "paragraph", Paragraph(child,document)
        elif child.tag==qn("w:tbl"): yield "table", Table(child,document)

def _docx(path:Path)->ExtractedText:
    from docx import Document
    doc=Document(str(path)); blocks=[]; output=[]; order=0; current_section="正文"
    for kind,item in _iter_docx_blocks(doc):
        if kind=="paragraph":
            text=item.text.strip()
            if not text:continue
            order+=1
            style=(item.style.name if item.style else "").lower()
            block_type="heading" if "heading" in style or "标题" in style or _HEADING_RE.match(text) else "paragraph"
            if block_type=="heading": current_section=text[:200]
            blocks.append(LayoutBlock(order,block_type,text,section=current_section))
            output.append(text)
        else:
            rows=[[_cell(c.text) for c in row.cells] for row in item.rows]
            markdown=table_markdown(rows)
            if markdown:
                order+=1; blocks.append(LayoutBlock(order,"table",markdown,section=current_section,table=rows)); output.append(markdown)
    text="\n\n".join(output)
    return ExtractedText(text=text,pages=[PageContent(1,text,_find_headings(text))],blocks=blocks,source_method="native_docx",parser_name="python-docx/OOXML-body-order")

def _xlsx(path:Path)->ExtractedText:
    from openpyxl import load_workbook
    wb=load_workbook(path,data_only=True,read_only=False); blocks=[]; output=[]; order=0
    for ws in wb.worksheets:
        order+=1; blocks.append(LayoutBlock(order,"heading",ws.title,section=ws.title,sheet_name=ws.title)); output.append(ws.title)
        rows=[[ws.cell(r,c).value for c in range(1,ws.max_column+1)] for r in range(1,ws.max_row+1)]
        for rng in ws.merged_cells.ranges:
            value=ws.cell(rng.min_row,rng.min_col).value
            for r in range(rng.min_row,rng.max_row+1):
                for c in range(rng.min_col,rng.max_col+1):
                    if rows[r-1][c-1] is None: rows[r-1][c-1]=value
        md=table_markdown(rows)
        if md:
            order+=1; blocks.append(LayoutBlock(order,"table",md,section=ws.title,table=[[ _cell(c) for c in row] for row in rows],sheet_name=ws.title)); output.append(md)
    text="\n\n".join(output)
    return ExtractedText(text=text,pages=[PageContent(1,text,_find_headings(text))],blocks=blocks,source_method="native_xlsx",parser_name="openpyxl/sheet-row-order")

def _pptx(path:Path)->ExtractedText:
    from pptx import Presentation
    prs=Presentation(str(path)); blocks=[]; output=[]; order=0
    for pno,slide in enumerate(prs.slides,1):
        order+=1; blocks.append(LayoutBlock(order,"heading",f"幻灯片 {pno}",page_number=pno,section=f"幻灯片 {pno}")); output.append(f"幻灯片 {pno}")
        shapes=sorted(enumerate(slide.shapes),key=lambda p:(int(getattr(p[1],"top",0)),int(getattr(p[1],"left",0)),p[0]))
        for _,shape in shapes:
            bbox=[float(shape.left),float(shape.top),float(shape.width),float(shape.height)]
            if getattr(shape,"has_table",False):
                rows=[[_cell(cell.text) for cell in row.cells] for row in shape.table.rows]; md=table_markdown(rows)
                if md:
                    order+=1; blocks.append(LayoutBlock(order,"table",md,pno,f"幻灯片 {pno}",bbox,rows)); output.append(md)
            elif getattr(shape,"has_text_frame",False) and shape.text.strip():
                order+=1; txt=shape.text.strip(); blocks.append(LayoutBlock(order,"paragraph",txt,pno,f"幻灯片 {pno}",bbox)); output.append(txt)
    text="\n\n".join(output)
    return ExtractedText(text=text,pages=[PageContent(1,text,_find_headings(text))],blocks=blocks,source_method="native_pptx",parser_name="python-pptx/visual-top-left-order")

def _pdf(path:Path)->ExtractedText:
    try: import fitz
    except Exception:
        return ExtractedText("",warning="未安装 PyMuPDF，无法进行 PDF 版面解析",source_method="native_pdf",parser_name="PyMuPDF")
    pdf=fitz.open(str(path)); blocks=[]; pages=[]; all_text=[]; order=0; scanned=[]
    for pno,page in enumerate(pdf,1):
        items=[]; table_rects=[]
        try:
            for table in page.find_tables().tables:
                rows=table.extract(); md=table_markdown(rows); rect=list(table.bbox)
                if md:
                    table_rects.append(rect); items.append((rect[1],rect[0],0,"table",md,rect,rows))
        except Exception: pass
        for x0,y0,x1,y1,text,*_ in page.get_text("blocks"):
            text=text.strip()
            if not text:continue
            if any(x0<r[2] and x1>r[0] and y0<r[3] and y1>r[1] for r in table_rects): continue
            items.append((y0,x0,1,"paragraph",text,[x0,y0,x1-x0,y1-y0],None))
        items.sort(key=lambda x:(x[0],x[1],x[2]))
        page_text=[]
        for _,_,_,kind,text,bbox,rows in items:
            order+=1; blocks.append(LayoutBlock(order,kind,text,pno,"正文",[float(v) for v in bbox],rows)); page_text.append(text)
        joined="\n".join(page_text); pages.append(PageContent(pno,joined,_find_headings(joined))); all_text.append(joined)
        if len(re.sub(r"\s+","",joined))<25: scanned.append(pno)
    return ExtractedText(text="\f".join(all_text),pages=pages,blocks=blocks,warning=(f"第 {', '.join(map(str,scanned))} 页疑似扫描件，需启用百度云端 OCR。" if scanned else ""),source_method="native_pdf",parser_name="PyMuPDF/page-y-x-order",requires_cloud_ocr=bool(scanned))

def _legacy(path:Path)->ExtractedText:
    suffix=path.suffix.lower(); target={".doc":"docx",".ppt":"pptx",".xls":"xlsx"}.get(suffix)
    command=shutil.which("soffice") or shutil.which("libreoffice")
    if not target or not command: return ExtractedText("",warning=f"{suffix} 旧版 Office 文件需安装 LibreOffice 转换为 OOXML，或使用百度云端文档解析。",source_method="legacy_office",parser_name="legacy-office")
    with tempfile.TemporaryDirectory(prefix="planner-office-") as dest:
        result=subprocess.run([command,"--headless","--convert-to",target,"--outdir",dest,str(path)],capture_output=True,text=True,timeout=120)
        converted=Path(dest)/f"{path.stem}.{target}"
        if result.returncode!=0 or not converted.exists(): return ExtractedText("",warning="旧版 Office 本地转换失败，请选择百度云端文档解析。",source_method="legacy_office",parser_name="LibreOffice")
        extracted=extract_text(converted, {"docx":"word","pptx":"ppt","xlsx":"excel"}[target])
        extracted.source_method="libreoffice_"+extracted.source_method; extracted.parser_name="LibreOffice → "+extracted.parser_name
        return extracted

def extract_text(path:str|Path,file_type:str="other")->ExtractedText:
    path=Path(path); suffix=path.suffix.lower()
    try:
        if suffix in {".txt",".md",".csv"}: return _plain(path)
        if suffix==".docx": return _docx(path)
        if suffix==".xlsx": return _xlsx(path)
        if suffix==".pptx": return _pptx(path)
        if suffix==".pdf": return _pdf(path)
        if suffix in {".doc",".xls",".ppt"}: return _legacy(path)
        if suffix in {".png",".jpg",".jpeg",".webp",".bmp",".tif",".tiff",".ofd"}:
            return ExtractedText("",warning="图片/扫描件不使用本地 OCR；请先配置并确认百度云端 OCR 风险后解析。",source_method="cloud_ocr_required",parser_name="none",requires_cloud_ocr=True)
        return _plain(path)
    except Exception as exc:
        return ExtractedText("",warning=f"解析失败：{exc}",source_method="native_unknown",parser_name="failed")
