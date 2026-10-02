from pathlib import Path
import httpx
import pytest
from app.services.text_extractor import extract_text
from app.adapters.search import SearchProvider


def test_excel_and_pptx_tables_are_structured_and_ordered(tmp_path: Path):
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    book = Workbook(); ws = book.active; ws.title='指标表'
    ws['A1']='规划指标'; ws.merge_cells('A1:B1'); ws.append(['常住人口','31400'])
    xlsx = tmp_path/'metrics.xlsx'; book.save(xlsx)
    x = extract_text(xlsx, 'excel')
    assert x.parser_name == 'openpyxl/sheet-row-order'
    assert len([b for b in x.blocks if b.block_type=='table']) == 1
    assert '31400' in x.blocks[-1].text

    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    early = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(0.4)); early.text_frame.paragraphs[0].text='先读标题'
    shape = slide.shapes.add_table(2,2,Inches(1),Inches(2),Inches(5),Inches(1))
    shape.table.cell(0,0).text='指标'; shape.table.cell(0,1).text='数值'; shape.table.cell(1,0).text='城镇化率'; shape.table.cell(1,1).text='76.9%'
    pptx=tmp_path/'layout.pptx'; presentation.save(pptx)
    p = extract_text(pptx, 'ppt')
    assert p.parser_name == 'python-pptx/visual-top-left-order'
    assert p.blocks[1].text == '先读标题'
    assert any(b.block_type=='table' and '76.9%' in b.text for b in p.blocks)


def test_pdf_reading_order_by_page_coordinates(tmp_path: Path):
    import fitz
    pdf = fitz.open(); page = pdf.new_page()
    page.insert_text((72, 72), 'TOP_BLOCK')
    page.insert_text((72, 180), 'BOTTOM_BLOCK')
    path = tmp_path/'layout.pdf'; pdf.save(path)
    result = extract_text(path, 'pdf')
    text_blocks=[b.text for b in result.blocks if b.block_type=='paragraph']
    assert text_blocks[:2] == ['TOP_BLOCK','BOTTOM_BLOCK']
    assert result.parser_name == 'PyMuPDF/page-y-x-order'


@pytest.mark.asyncio
async def test_configured_search_adapter_uses_real_provider_request(monkeypatch):
    seen = {}
    class FakeClient:
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return False
        async def post(self, url, headers=None, json=None, **kwargs):
            seen.update(url=url, headers=headers, json=json)
            return httpx.Response(200, json={'results':[{'title':'规划政策','url':'https://example.gov.cn/policy','content':'政策摘要'}]})
    monkeypatch.setattr('app.adapters.search.httpx.AsyncClient', lambda **kwargs: FakeClient())
    rows = await SearchProvider.search('国土空间规划', provider_config={'search_provider':'Tavily','search_api_key':'demo-key'})
    assert seen['url'] == 'https://api.tavily.com/search'
    assert seen['headers']['Authorization'] == 'Bearer demo-key'
    assert seen['json']['query'] == '国土空间规划'
    assert rows[0]['title'] == '规划政策'
